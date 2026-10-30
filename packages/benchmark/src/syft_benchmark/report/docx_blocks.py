"""The python-docx building blocks shared by the report documents."""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence

from docx.document import Document as Doc
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

TABLE_STYLE = "Light Grid Accent 1"
MUTED = RGBColor(0x55, 0x55, 0x55)


def heading(doc: Doc, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Doc, text: str, *, bold: bool = False, italic: bool = False) -> None:
    run = doc.add_paragraph().add_run(text)
    run.bold = bold
    run.italic = italic


def note(doc: Doc, text: str) -> None:
    """An explanation under a section: what the quantity means and what it is not."""
    run = doc.add_paragraph().add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def image(doc: Doc, png: bytes, width: float = 6.0) -> None:
    doc.add_picture(io.BytesIO(png), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def table(doc: Doc, headers: Sequence[str], rows: Iterable[Sequence[str]]) -> None:
    grid = doc.add_table(rows=1, cols=len(headers), style=TABLE_STYLE)
    for cell, name in zip(grid.rows[0].cells, headers, strict=True):
        cell.text = name
    for row in rows:
        cells = grid.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value
    for row in grid.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()
