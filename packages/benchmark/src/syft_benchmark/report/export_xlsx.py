"""The Excel export of one run (job): everything that took part, full texts.

Eight sheets (report API, "Excel export"): Run, Generated, Rejected, Kept,
Filter, Answers, Judging, Timing. Rows follow the question order
(``question_order``). A text longer than an Excel cell holds continues in
"(part 2)", "(part 3)" columns of the same row. What a run did not record is
left empty.

Written with openpyxl in write-only mode: one sheet's rows are built in memory
(the part columns must be known before the header), written, then dropped.
"""

from __future__ import annotations

import json
import math
import re
from collections import defaultdict
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import IO, Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from syft_benchmark.config import PairStatus, StatusReason
from syft_benchmark.db.models import Job, QaPair, Result, Run, Target
from syft_benchmark.generation import decisions, recorded
from syft_benchmark.generation.article_dates import DATE_KEY
from syft_benchmark.question_order import answer_order, pair_key, pair_order
from syft_benchmark.report import run_view
from syft_benchmark.report.filter_view import _inferred, _rows
from syft_benchmark.runs.judge_stage import OWNER_OVERRIDE, behavior_of

# Excel's limit on the characters of one cell.
CELL_LIMIT = 32_767

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SHEETS = (
    "Run",
    "Generated",
    "Rejected",
    "Kept",
    "Filter",
    "Answers",
    "Judging",
    "Timing",
)

TEXT_METRICS = ("bleu", "rouge1_f", "rouge2_f", "rougeL_f", "bertscore_f1")

# The Rejected stage of questions the set cap took out.
CAP = "cap"


def filename(endpoint: str, created_at: datetime | None) -> str:
    """``<endpoint>-run-<YYYY-MM-DD-HHMM>.xlsx``."""
    clean = re.sub(r"[^A-Za-z0-9]+", "-", endpoint).strip("-").lower() or "endpoint"
    stamp = (
        created_at.astimezone(UTC).strftime("%Y-%m-%d-%H%M")
        if created_at
        else "undated"
    )
    return f"{clean}-run-{stamp}.xlsx"


# --- cells ------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class Col:
    title: str
    width: float = 14
    # Free text: wrapped. Any column splits into parts past the cell limit.
    long: bool = False


def _text(value: Any) -> str:
    """A cell-safe string (no control characters Excel refuses)."""
    return str(ILLEGAL_CHARACTERS_RE.sub("", str(value)))


def _value(value: Any) -> Any:
    """A Python value as a cell value; None stays empty."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, int | float):
        return value
    if isinstance(value, datetime):
        return value.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dict | list):
        return _text(json.dumps(value, ensure_ascii=False, default=str))
    return _text(value)


def _parts(text: str) -> list[str]:
    return [text[i : i + CELL_LIMIT] for i in range(0, len(text), CELL_LIMIT)]


class _Styles:
    """The sheet's cell styles, resolved once: setting a style per cell
    costs more than writing it."""

    def __init__(self, ws: WriteOnlyWorksheet) -> None:
        def resolved(**style: Any) -> Any:
            cell = WriteOnlyCell(ws)
            for name, value in style.items():
                setattr(cell, name, value)
            return cell._style

        self.bold = resolved(font=Font(bold=True))
        self.wrap = resolved(alignment=Alignment(wrap_text=True, vertical="top"))
        self.top = resolved(alignment=Alignment(vertical="top"))


def _cell(ws: WriteOnlyWorksheet, value: Any, style: Any) -> Any:
    if value is None or isinstance(value, int | float):
        return value
    cell = WriteOnlyCell(ws, value)
    # A text starting with "=" is text, not a formula.
    cell.data_type = "s"
    cell._style = style
    return cell


def _sheet(
    wb: Workbook,
    name: str,
    cols: Sequence[Col],
    rows: Sequence[Sequence[Any]],
) -> None:
    """One sheet: header, part columns where a text runs past the limit."""
    cells = [[_value(v) for v in row] for row in rows]
    parts = [1] * len(cols)
    for row in cells:
        for i, value in enumerate(row):
            if isinstance(value, str) and len(value) > CELL_LIMIT:
                parts[i] = max(parts[i], math.ceil(len(value) / CELL_LIMIT))

    headers: list[str] = []
    widths: list[float] = []
    wraps: list[bool] = []
    for col, n in zip(cols, parts, strict=True):
        for k in range(1, n + 1):
            headers.append(col.title if k == 1 else f"{col.title} (part {k})")
            widths.append(col.width)
            wraps.append(col.long)

    ws = wb.create_sheet(name)
    style = _Styles(ws)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(cells) + 1}"
    head = []
    for title in headers:
        cell = WriteOnlyCell(ws, title)
        cell.data_type = "s"
        cell._style = style.bold
        head.append(cell)
    ws.append(head)
    for row in cells:
        out: list[Any] = []
        for value, n in zip(row, parts, strict=True):
            if n == 1:
                out.append(value)
                continue
            chunks: list[Any] = _parts(value) if isinstance(value, str) else [value]
            out.extend([*chunks, *([None] * (n - len(chunks)))])
        ws.append(
            [
                _cell(ws, v, style.wrap if wraps[i] else style.top)
                for i, v in enumerate(out)
            ]
        )


# --- reading what was recorded --------------------------------------------


def _dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _first(record: dict[str, Any], *keys: str) -> Any:
    """The first non-empty value under ``keys``."""
    for key in keys:
        value = record.get(key)
        if value not in (None, "", [], {}):
            return value
    return None


def _lines(items: Iterable[Any]) -> str:
    return "\n".join(f"- {item}" for item in items)


def _citations(found: Any) -> str:
    rows = []
    for item in found if isinstance(found, list) else []:
        if isinstance(item, dict) and item.get("url"):
            title = str(item.get("title") or "")
            rows.append(f"{item['url']} {title}".strip())
    return "\n".join(rows)


def _seconds(start: datetime | None, end: datetime | None) -> float | None:
    if start is None or end is None:
        return None
    return round((end - start).total_seconds(), 2)


# --- questions --------------------------------------------------------------

PAIR_COLS = (
    Col("Question id", 34),
    Col("Kind", 22),
    Col("Task type", 12),
    Col("Article", 30, long=True),
    Col("File", 24),
    Col("Article date", 12),
    Col("Passage", 60, long=True),
    Col("Writer model", 26),
    Col("Masked by", 24),
    Col("Writer prompt (system)", 60, long=True),
    Col("Writer prompt (user)", 60, long=True),
    Col("Writer reply", 60, long=True),
    Col("Writer call", 34),
    Col("Writer call time (s)", 10),
    Col("Writer call cost (USD)", 10),
    Col("Question", 50, long=True),
    Col("Right answer", 40, long=True),
    Col("Options / key facts / blanks", 40, long=True),
    Col("Expected behaviour", 12),
    Col("Status", 10),
    Col("Status reason", 14),
    Col("Status note", 30, long=True),
    Col("Written by job", 34),
    Col("Written at (UTC)", 20),
)


Calls = dict[str, dict[str, Any]]


def _call_id(row: QaPair) -> str:
    writer = (row.meta or {}).get(recorded.WRITER)
    return str(writer.get("call") or "") if isinstance(writer, dict) else ""


def _writer_calls(session: Session, rows: Sequence[QaPair]) -> Calls:
    """The writer calls behind ``rows`` (``recorded.writer_calls``). A call's
    texts sit on one of its pairs only; one not among ``rows`` is read from
    the database."""
    calls = recorded.writer_calls(rows)
    missing = sorted({_call_id(row) for row in rows} - set(calls) - {""})
    if missing:
        column = QaPair.meta[recorded.WRITER]["call"].astext
        found = session.scalars(select(QaPair).where(column.in_(missing))).all()
        calls.update(recorded.writer_calls(found))
    return calls


def _details(row: QaPair) -> str:
    """Options, key facts, blanks: whatever the kind carries."""
    meta = row.meta or {}
    out: list[str] = []
    if row.distractors:
        out.append("Other options:\n" + _lines(row.distractors))
    facts = _first(meta, "key_facts", "hop_facts")
    if isinstance(facts, list):
        out.append("Key facts:\n" + _lines(facts))
    claims = meta.get("claims")
    if isinstance(claims, list) and claims and claims != facts:
        out.append("Claims checked:\n" + _lines(claims))
    if meta.get("category") or meta.get("entity_label"):
        label = " ".join(
            str(meta[k]) for k in ("category", "entity_label") if meta.get(k)
        )
        out.append(f"Blank: {label}")
    for key, title in (("premise", "Premise"), ("missing", "Missing")):
        if meta.get(key):
            out.append(f"{title}: {meta[key]}")
    return "\n\n".join(out)


def _masked_by(meta: dict[str, Any]) -> str:
    """Masking kinds: "spaCy <model>" or "LLM: <why>"; empty for other kinds."""
    if meta.get("mode") == "spacy":
        return " ".join(["spaCy", str(meta.get("spacy_model") or "")]).strip()
    if meta.get("mode") == "llm":
        return f"LLM: {meta['llm_why']}" if meta.get("llm_why") else "LLM"
    return ""


def _pair_cells(row: QaPair, calls: Calls) -> list[Any]:
    call_id = _call_id(row)
    call = calls.get(call_id, {})
    return [
        row.id,
        row.generator,
        row.task_type,
        row.document_title,
        row.file_name,
        (row.meta or {}).get(DATE_KEY),
        call.get("passage") or row.context,
        call.get("model") or row.model,
        _masked_by(row.meta or {}),
        call.get("system"),
        call.get("user"),
        call.get("reply"),
        call_id,
        call.get("latency_s"),
        call.get("cost_usd"),
        row.question,
        row.answer,
        _details(row),
        row.expected_behavior,
        row.status,
        row.status_reason,
        row.status_note,
        row.job_id,
        row.created_at,
    ]


# --- the job's filter decisions --------------------------------------------


@dataclass(slots=True)
class _Decision:
    row: QaPair
    entry: dict[str, Any]
    web: dict[str, Any]
    recorded: bool


def _decisions(rows: Sequence[QaPair], job_id: str) -> list[_Decision]:
    """Every filter decision of the job, as ``filter_view`` lists them."""
    out: list[_Decision] = []
    for row in rows:
        meta = row.meta or {}
        entry = decisions.of_job(meta, job_id)
        recorded = entry is not None
        if entry is None:
            if row.job_id != job_id or meta.get(decisions.SCREENING):
                continue
            entry = _inferred(row)
            if entry is None:
                continue
            web = _dict(meta.get("web_check"))
        else:
            web = _dict(entry.get("web"))
        out.append(_Decision(row, entry, web, recorded))
    return out


FILTER_COLS = (
    Col("Question id", 34),
    Col("Kind", 22),
    Col("Question", 50, long=True),
    Col("Right answer", 40, long=True),
    Col("Stage", 12),
    Col("Outcome", 10),
    Col("Reason", 16),
    Col("Note", 40, long=True),
    Col("Decided at (UTC)", 20),
    Col("Recorded", 9),
    Col("Written by job", 34),
    Col("Written earlier", 9),
    Col("Status now", 10),
    Col("Control gate model", 26),
    Col("Control gate prompt (system)", 60, long=True),
    Col("Control gate prompt (user)", 60, long=True),
    Col("Control gate reply", 60, long=True),
    Col("Web check model", 26),
    Col("Web check prompt (system)", 60, long=True),
    Col("Web check prompt (user)", 50, long=True),
    Col("Web check answer", 60, long=True),
    Col("Citations", 50, long=True),
    Col("Searches", 9),
    Col("Engine", 9),
    Col("Searched", 9),
    Col("Judge", 26),
    Col("Judge prompt (system)", 60, long=True),
    Col("Judge prompt (user)", 60, long=True),
    Col("Judge reply", 60, long=True),
    Col("Verdict", 12),
    Col("Reasoning", 50, long=True),
    Col("Error", 30, long=True),
    Col("Time (s)", 9),
    Col("Cost (USD)", 10),
)


def _total(*values: Any) -> float | None:
    """The sum of the numbers given; None when there is none."""
    found = [
        v for v in values if isinstance(v, int | float) and not isinstance(v, bool)
    ]
    return round(sum(found), 6) if found else None


def _filter_cells(d: _Decision, job_id: str) -> list[Any]:
    row, entry, web = d.row, d.entry, d.web
    gate = _dict(entry.get("gate"))
    requests = web.get("web_search_requests")
    return [
        row.id,
        row.generator,
        row.question,
        row.answer,
        entry.get("stage"),
        entry.get("outcome"),
        entry.get("reason"),
        entry.get("note"),
        entry.get("at"),
        d.recorded,
        row.job_id,
        row.job_id != job_id,
        row.status,
        gate.get("model"),
        gate.get("system"),
        gate.get("user"),
        gate.get("reply"),
        web.get("model"),
        web.get("system"),
        web.get("user"),
        web.get("answer"),
        _citations(web.get("citations")),
        requests if isinstance(requests, int) else None,
        web.get("search"),
        web.get("web_search") if "web_search" in web else None,
        web.get("judge"),
        web.get("judge_system"),
        web.get("judge_prompt"),
        web.get("judge_reply"),
        web.get("verdict"),
        web.get("reasoning"),
        web.get("error"),
        _total(gate.get("latency_s"), web.get("latency_s"), web.get("judge_latency_s")),
        _total(gate.get("cost_usd"), web.get("cost_usd"), web.get("judge_cost_usd")),
    ]


# --- answers and judging ----------------------------------------------------


ANSWER_COLS = (
    Col("Question id", 34),
    Col("Kind", 22),
    Col("Question", 40, long=True),
    Col("Model", 26),
    Col("Condition", 18),
    Col("Check", 22),
    Col("For judge", 22),
    Col("System prompt", 50, long=True),
    Col("User prompt", 60, long=True),
    Col("Excerpts sent", 60, long=True),
    Col("Answer", 60, long=True),
    Col("Web search", 10),
    Col("Searches", 9),
    Col("Citations", 40, long=True),
    Col("Search unused", 9),
    Col("Finish reason", 10),
    Col("Truncated", 9),
    Col("Reused", 9),
    Col("Served by", 14),
    Col("Retrieval hit", 9),
    Col("Retrieval rank", 9),
    Col("Latency (s)", 9),
    Col("Cost (USD)", 10),
    *(Col(name, 10) for name in TEXT_METRICS),
)
_A = {col.title: n for n, col in enumerate(ANSWER_COLS)}

JUDGING_COLS = (
    Col("Judge", 26),
    Col("Question id", 34),
    Col("Kind", 22),
    Col("Question", 40, long=True),
    Col("Model", 26),
    Col("Condition", 18),
    Col("Check", 22),
    Col("Judge prompt (system)", 60, long=True),
    Col("Judge prompt (user)", 60, long=True),
    Col("Judge reply", 60, long=True),
    Col("Verdict", 12),
    Col("Behaviour", 12),
    Col("Reasoning", 50, long=True),
    Col("Judged without a model", 9),
    Col("Served by", 14),
    Col("Latency (s)", 9),
    Col("Cost (USD)", 10),
    Col("Primary judge", 9),
    Col("Counted", 9),
    Col("Judged at (UTC)", 20),
)
_J = {col.title: n for n, col in enumerate(JUDGING_COLS)}


@dataclass(slots=True)
class _Answer:
    result: Result
    run: Run
    generator: str
    question: str


def _job_answers(session: Session, job_id: str) -> list[_Answer]:
    """Every result row of the job, in the question order."""
    stmt = (
        select(Result, Run, QaPair.generator, QaPair.question)
        .join(Run, Run.id == Result.run_id)
        .outerjoin(QaPair, QaPair.id == Result.qa_id)
        .where(Run.job_id == job_id)
        .order_by(
            *pair_order(QaPair.generator, QaPair.created_at, Result.qa_id),
            *answer_order(
                Run.context_mode,
                Run.block,
                Run.model,
                Result.judge_model,
                Result.created_at,
                Result.id,
            ),
        )
    )
    return [
        _Answer(result, run, generator or "", question or "")
        for result, run, generator, question in session.execute(stmt)
    ]


def _grouped(rows: Sequence[_Answer]) -> dict[tuple[str, ...], list[_Answer]]:
    """The rows of one answer (model x arm x block x question) together."""
    groups: dict[tuple[str, ...], list[_Answer]] = defaultdict(list)
    for a in rows:
        key = (
            a.run.context_mode,
            a.run.context_source,
            a.run.block,
            a.run.model,
            a.result.qa_id,
        )
        groups[key].append(a)
    return groups


def _log(block: dict[str, Any]) -> list[dict[str, Any]]:
    log = block.get("log")
    return [e for e in log if isinstance(e, dict)] if isinstance(log, list) else []


def _round(entry: dict[str, Any]) -> str:
    return f"challenge round {entry.get('round')}"


def _trial(entry: dict[str, Any]) -> str:
    return (
        f"repeat at temperature {entry.get('temperature')}, trial {entry.get('trial')}"
    )


def _answer_rows(rows: Sequence[_Answer]) -> list[list[Any]]:
    """One answer, then the challenge rounds and repeats of each judge's
    block (every judge seat drives its own)."""
    lead = rows[0]
    result, run = lead.result, lead.run
    audit = _dict(result.audit)
    call = _dict(audit.get("call"))
    metrics = _dict(_dict(result.extra).get("text_metrics"))
    requests = call.get("web_search_requests")
    head = [result.qa_id, lead.generator, lead.question, run.model, run.context_mode]
    first = [
        *head,
        run.block if run.block == "direct" else f"{run.block}: first answer",
        None,
        audit.get("responder_system"),
        audit.get("responder_prompt"),
        audit.get("context"),
        result.answer,
        call.get("web_search") if "web_search" in call else None,
        requests if isinstance(requests, int) else None,
        _citations(audit.get("citations")),
        audit.get("web_search_unused") if "web_search_unused" in audit else None,
        call.get("finish_reason"),
        call.get("truncated") if "truncated" in call else None,
        call.get("reused") if "reused" in call else None,
        result.served_by,
        result.retrieval_hit,
        result.retrieval_rank,
        result.latency_s,
        call.get("cost_usd"),
        *(metrics.get(name) for name in TEXT_METRICS),
    ]
    out = [first]

    def sub(check: str, judge: str, entry: dict[str, Any]) -> list[Any]:
        row: list[Any] = [*head, *[None] * (len(ANSWER_COLS) - len(head))]
        row[_A["Check"]] = check
        row[_A["For judge"]] = judge
        row[_A["Answer"]] = entry.get("answer")
        row[_A["Latency (s)"]] = entry.get("latency_s")
        row[_A["Cost (USD)"]] = entry.get("cost_usd")
        return row

    for a in rows:
        extra = _dict(a.result.extra)
        own = _dict(a.result.audit)
        for entry in _log(_dict(extra.get("denial"))):
            row = sub(_round(entry), a.result.judge_model, entry)
            row[_A["User prompt"]] = entry.get("objection")
            out.append(row)
        for entry in _log(_dict(extra.get("monte_carlo"))):
            row = sub(_trial(entry), a.result.judge_model, entry)
            row[_A["System prompt"]] = own.get("responder_system")
            row[_A["User prompt"]] = own.get("responder_prompt")
            out.append(row)
    return out


def _judging_rows(a: _Answer, primary: str | None, counted: bool) -> list[list[Any]]:
    """One judge's verdict on an answer, then its calls on the block's rounds
    and repeats."""
    result, run = a.result, a.run
    audit = _dict(result.audit)
    judge = _dict(audit.get("judge_call"))
    is_primary = result.judge_model == primary if primary else None
    main = [
        result.judge_model,
        result.qa_id,
        a.generator,
        a.question,
        run.model,
        run.context_mode,
        run.block,
        audit.get("judge_system"),
        audit.get("judge_prompt"),
        audit.get("judge_raw"),
        result.verdict,
        behavior_of(result.extra),
        result.reasoning,
        audit.get("judged_without_model") if audit else None,
        result.judge_served_by,
        judge.get("latency_s"),
        judge.get("cost_usd"),
        is_primary,
        counted,
        result.created_at,
    ]
    out = [main]
    extra = _dict(result.extra)
    for name, label in (("denial", _round), ("monte_carlo", _trial)):
        block = _dict(extra.get(name))
        for entry in _log(block):
            if not (entry.get("judge_prompt") or entry.get("judge_raw")):
                continue
            row: list[Any] = [*main[:6], *[None] * (len(JUDGING_COLS) - 6)]
            row[_J["Check"]] = label(entry)
            row[_J["Judge prompt (system)"]] = block.get("judge_system")
            row[_J["Judge prompt (user)"]] = entry.get("judge_prompt")
            row[_J["Judge reply"]] = entry.get("judge_raw")
            if "correct" in entry:
                row[_J["Verdict"]] = "correct" if entry["correct"] else "not correct"
            row[_J["Latency (s)"]] = entry.get("judge_latency_s")
            row[_J["Cost (USD)"]] = entry.get("judge_cost_usd")
            row[_J["Primary judge"]] = is_primary
            out.append(row)
    return out


# --- the Run sheet ----------------------------------------------------------


def _flat(prefix: str, value: Any, out: list[list[Any]]) -> None:
    if isinstance(value, dict) and value:
        for key, inner in value.items():
            _flat(f"{prefix}.{key}" if prefix else str(key), inner, out)
    else:
        out.append([prefix, value])


def _run_rows(
    session: Session,
    job: Job,
    target: Target,
    report: dict[str, Any],
    counts: dict[str, int],
    asked: int,
) -> list[list[Any]]:
    params = dict(job.params or {})
    run = _dict(report.get("run"))
    method = _dict(report.get("method"))
    funnel = _dict(report.get("funnel"))
    cost = _dict(params.get(run_view.COST))
    timing = _dict(params.get(run_view.TIMING))
    rows: list[list[Any]] = [
        ["Endpoint", target.endpoint or target.key],
        ["Target", target.key],
        ["Job", job.id],
        ["Job kind", job.kind],
        ["Trigger", job.trigger],
        ["State", job.state],
        ["Error", job.error],
        ["Created (UTC)", job.created_at],
        ["Started (UTC)", job.started_at],
        ["Finished (UTC)", job.finished_at],
        ["Duration (s)", _seconds(job.started_at, job.finished_at)],
        ["Window (days)", run.get("window_days")],
        ["Window from", method.get("articles_from")],
        ["Window to", method.get("articles_to")],
        ["Articles", run.get("articles")],
        ["Articles dated", run.get("articles_dated")],
        ["Articles published from", run.get("articles_first")],
        ["Articles published to", run.get("articles_last")],
        ["Articles new this run", run.get("articles_new")],
        ["Written", counts.get("written")],
        ["Filter decisions", counts.get("checked")],
        ["Kept", counts.get("kept")],
        ["Rejected", counts.get("removed")],
        ["Filter failed", counts.get("failed")],
        ["Filter skipped", counts.get("skipped")],
        ["Asked", asked],
        ["Questions in the figures", run.get("questions")],
        ["Cost total (USD)", cost.get("total_usd")],
    ]
    by_role = _dict(cost.get("by_role"))
    for role in ("writer", "web_check", "subjects", "judges"):
        rows.append([f"Cost: {role} (USD)", by_role.get(role)])
    rows += [
        ["Cost of all priced calls (USD)", cost.get("usd_calls")],
        ["Key spend before (USD)", cost.get("spend_before")],
        ["Key spend after (USD)", cost.get("spend_after")],
        ["Key spend delta (USD)", cost.get("usd")],
        ["Timing total (s)", timing.get("total_s")],
    ]
    for phase in timing.get("phases") or []:
        if isinstance(phase, dict):
            rows.append([f"Phase: {phase.get('phase')} (s)", phase.get("s")])
    models = sorted(
        set(
            session.scalars(
                select(Run.model).where(Run.job_id == job.id, Run.model != "")
            )
        )
    )
    rows += [
        ["Tested models", "\n".join(models)],
        ["Judges", "\n".join(str(j) for j in method.get("judges") or [])],
        ["Judge panel (launch)", "\n".join(run_view.job_panel(job))],
        ["Primary judge", params.get(run_view.PRIMARY_JUDGE)],
        ["Judge policy", params.get(run_view.JUDGE_POLICY)],
        ["Web check model", params.get(run_view.WEB_CHECK_MODEL)],
        ["Web check judge", params.get(run_view.WEB_CHECK_JUDGE)],
        ["Question writer", method.get("generator_model")],
        ["Kinds", "\n".join(str(k) for k in method.get("kinds") or [])],
    ]
    concurrency = _dict(timing.get("concurrency"))
    rows += [
        ["Concurrency: model", concurrency.get("model")],
        ["Concurrency: endpoint", concurrency.get("endpoint")],
        ["Profile", method.get("profile")],
        ["Excerpts per question", method.get("context_docs")],
        ["Challenge rounds", method.get("denial_rounds")],
        ["Repeats", method.get("repeats")],
    ]
    for key, value in funnel.items():
        rows.append([f"Funnel: {key}", value])

    # Everything else the launch kept: the request and settings snapshot.
    shown = {
        run_view.COST,
        run_view.TIMING,
        run_view.GENERATION,
        run_view.PRIMARY_JUDGE,
        run_view.JUDGE_POLICY,
        run_view.WEB_CHECK_MODEL,
        run_view.WEB_CHECK_JUDGE,
        run_view.JUDGE_PANEL,
    }
    for key, value in params.items():
        if key not in shown:
            _flat(f"Launch: {key}", value, rows)
    for kind in run_view.generation_of(job):
        _flat(f"Build: {kind.get('kind')}", kind, rows)

    # The methodology snapshot of the job's runs; a value that differed
    # between runs lists each.
    seen: dict[str, list[Any]] = {}
    for run_params in session.scalars(
        select(Run.params).where(
            Run.job_id == job.id, Run.judge_model != OWNER_OVERRIDE
        )
    ):
        for key, value in _dict(run_params).items():
            values = seen.setdefault(str(key), [])
            if value not in values:
                values.append(value)
    for key in sorted(seen):
        values = seen[key]
        rows.append([f"Settings: {key}", values[0] if len(values) == 1 else values])
    return rows


def _timing_rows(job: Job) -> list[list[Any]]:
    timing = _dict((job.params or {}).get(run_view.TIMING))
    rows: list[list[Any]] = []
    if timing.get("total_s") is not None:
        rows.append(["total", *[None] * 5, timing.get("total_s")])
    for phase in timing.get("phases") or []:
        if isinstance(phase, dict):
            rows.append(
                ["phase", phase.get("phase"), None, None, None, None, phase.get("s")]
            )
    for p in timing.get("passes") or []:
        if isinstance(p, dict):
            rows.append(
                [
                    "pass",
                    None,
                    p.get("arm"),
                    p.get("block"),
                    p.get("model"),
                    None,
                    p.get("s"),
                    p.get("questions"),
                    p.get("stopped"),
                ]
            )
    for c in timing.get("calls") or []:
        if isinstance(c, dict):
            rows.append(
                [
                    "calls",
                    None,
                    None,
                    None,
                    c.get("model"),
                    c.get("role"),
                    c.get("total_s"),
                    None,
                    None,
                    c.get("count"),
                    c.get("failed"),
                    c.get("mean_s"),
                    c.get("p90_s"),
                    c.get("max_s"),
                ]
            )
    width = len(TIMING_COLS)
    return [[*row, *[None] * (width - len(row))] for row in rows]


TIMING_COLS = (
    Col("Part", 8),
    Col("Phase", 10),
    Col("Condition", 18),
    Col("Check", 12),
    Col("Model", 26),
    Col("Role", 8),
    Col("Seconds", 10),
    Col("Questions", 10),
    Col("Stopped", 20),
    Col("Calls", 8),
    Col("Failed", 8),
    Col("Mean (s)", 9),
    Col("p90 (s)", 9),
    Col("Max (s)", 9),
)


# --- the workbook -----------------------------------------------------------


def write(
    session: Session,
    job: Job,
    target: Target,
    out: IO[bytes] | str,
    *,
    configured: run_view.Panel = (),
) -> None:
    """Write the run's workbook to ``out`` (a path or a binary file)."""
    wb = Workbook(write_only=True)
    report = run_view.report_part(session, job, configured=configured)
    primary = run_view.primary_judge(session, job, configured=configured)
    judges = run_view.judge_order(
        (), primary, run_view.job_panel(job) or _resolved(configured)
    )

    pairs = _rows(session, job.target, [job.id])
    found = _decisions(pairs, job.id)
    counts = {
        "written": sum(1 for row in pairs if row.job_id == job.id),
        "checked": len(found),
        **decisions.outcome_counts([d.entry for d in found]),
    }
    answers = _job_answers(session, job.id)
    asked_ids = list(dict.fromkeys(a.result.qa_id for a in answers))
    judges = run_view.judge_order(
        (a.result.judge_model for a in answers), primary, judges
    )

    _sheet(
        wb,
        "Run",
        (Col("Field", 34), Col("Value", 60, long=True)),
        _run_rows(session, job, target, report, counts, len(asked_ids)),
    )

    written = _ordered(row for row in pairs if row.job_id == job.id)
    kept = _kept(session, job, pairs, found, asked_ids)
    calls = _writer_calls(session, [*pairs, *kept])
    _sheet(wb, "Generated", PAIR_COLS, [_pair_cells(r, calls) for r in written])
    _sheet(
        wb,
        "Rejected",
        (Col("Stage", 14), Col("Reason", 16), Col("Note", 40, long=True), *PAIR_COLS),
        _rejected(written, found, calls),
    )
    _sheet(wb, "Kept", PAIR_COLS, [_pair_cells(r, calls) for r in kept])

    found.sort(key=lambda d: _key(d.row))
    _sheet(wb, "Filter", FILTER_COLS, [_filter_cells(d, job.id) for d in found])
    del pairs, found, written, kept, calls

    groups = _grouped(answers)
    answer_rows: list[list[Any]] = []
    for rows in groups.values():
        graded = [a for a in rows if a.result.judge_model != OWNER_OVERRIDE]
        answer_rows.extend(_answer_rows(graded or rows))
    _sheet(wb, "Answers", ANSWER_COLS, answer_rows)
    del answer_rows

    panel = [*judges, ""]
    judging: list[list[Any]] = []
    for rows in groups.values():
        chosen = run_view.pick([a.result for a in rows], panel).counted
        for a in rows:
            if a.result.judge_model == "":
                continue
            judging.extend(_judging_rows(a, primary, a.result is chosen))
    _sheet(wb, "Judging", JUDGING_COLS, judging)
    del judging, groups, answers

    _sheet(wb, "Timing", TIMING_COLS, _timing_rows(job))
    wb.save(out)


def _resolved(configured: run_view.Panel) -> list[str]:
    return list(configured() if callable(configured) else configured)


def _key(row: QaPair) -> tuple[Any, ...]:
    return pair_key(row.generator, row.created_at, row.id)


def _ordered(rows: Iterable[QaPair]) -> list[QaPair]:
    return sorted(rows, key=_key)


def _rejected(
    written: Sequence[QaPair], found: Sequence[_Decision], calls: Calls
) -> list[list[Any]]:
    """The job's removals: its filter decisions, and its questions the set cap
    took out."""
    out: list[tuple[tuple[Any, ...], list[Any]]] = []
    removed = set()
    for d in found:
        if d.entry.get("outcome") != decisions.REMOVED:
            continue
        removed.add(d.row.id)
        entry = d.entry
        cells = [entry.get("stage"), entry.get("reason"), entry.get("note")]
        out.append((_key(d.row), [*cells, *_pair_cells(d.row, calls)]))
    for row in written:
        if row.id in removed or row.status_reason != StatusReason.ROTATION.value:
            continue
        if row.status not in (PairStatus.REJECTED.value, PairStatus.RETIRED.value):
            continue
        cells = [CAP, row.status_reason, row.status_note]
        out.append((_key(row), [*cells, *_pair_cells(row, calls)]))
    out.sort(key=lambda item: item[0])
    return [cells for _, cells in out]


def _kept(
    session: Session,
    job: Job,
    pairs: Sequence[QaPair],
    found: Sequence[_Decision],
    asked_ids: Sequence[str],
) -> list[QaPair]:
    """The questions the run asked; for a build-only job, the ones it let
    into the set."""
    if asked_ids:
        return _ordered(
            session.scalars(select(QaPair).where(QaPair.id.in_(set(asked_ids))))
        )
    kept_ids = {d.row.id for d in found if d.entry.get("outcome") == decisions.KEPT}
    return _ordered(
        r
        for r in pairs
        if r.status == PairStatus.ACTIVE.value
        and (r.id in kept_ids or r.job_id == job.id)
    )
