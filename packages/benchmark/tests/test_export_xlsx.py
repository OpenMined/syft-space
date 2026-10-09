"""The Excel export of a run: sheets, rows in the question order, full texts."""

from __future__ import annotations

import io
import time
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete

from syft_benchmark.config import JobState, PairStatus, StatusReason, get_settings
from syft_benchmark.control.app import create_app
from syft_benchmark.control.schemas import TargetSpec
from syft_benchmark.db import Job, QaPair, Result, Run, Target, session_scope
from syft_benchmark.generation import decisions
from syft_benchmark.generation.recorded import PASSAGE_MARK
from syft_benchmark.report import export_xlsx

TOKEN = "test-export-token"
KEY = "pytest-export-a"
KEY2 = "pytest-export-b"
AUTH = {"Authorization": f"Bearer {TOKEN}"}
JOB = "xl-job"
EARLIER = "xl-earlier"
T0 = datetime(2026, 9, 1, 6, 0, tzinfo=UTC)
M1, M2 = "vendor/model-one", "other/model-two"
J1, J2 = "judge/primary", "judge/second"
# Over two cell limits: three columns.
LONG = "".join(f"{n:08d} " for n in range(8_000))


def _db_ready() -> bool:
    try:
        with session_scope() as session:
            session.execute(delete(QaPair).where(QaPair.space == KEY))
        return True
    except Exception:  # noqa: BLE001 - the database may simply not be at hand
        return False


pytestmark = pytest.mark.skipif(
    not _db_ready(), reason="the benchmark database is unavailable"
)


@pytest.fixture
def clean() -> Any:
    def wipe() -> None:
        with session_scope() as session:
            for key in (KEY, KEY2):
                session.execute(delete(Result).where(Result.space == key))
                session.execute(delete(Run).where(Run.space == key))
                session.execute(delete(QaPair).where(QaPair.space == key))
                session.execute(delete(Job).where(Job.target == key))
                session.execute(delete(Target).where(Target.key == key))

    wipe()
    yield
    wipe()


@pytest.fixture
def client() -> TestClient:
    conf = get_settings().model_copy(update={"control_token": TOKEN})
    return TestClient(create_app(conf))


def _console(client: TestClient, key: str = KEY) -> dict[str, str]:
    body = TargetSpec(key=key, url="http://space.invalid", endpoint="Support KB")
    put = client.put(f"/targets/{key}", json=body.model_dump(mode="json"), headers=AUTH)
    assert put.status_code == 200, put.text
    minted = client.post(f"/targets/{key}/session", headers=AUTH)
    assert minted.status_code == 200, minted.text
    return {"Authorization": f"Bearer {minted.json()['token']}"}


def _pair(qa: str, generator: str, minute: int, **extra: Any) -> QaPair:
    return QaPair(
        id=qa,
        space=KEY,
        job_id=extra.pop("job_id", JOB),
        generator=generator,
        document_title=f"Article {qa}",
        file_name=f"{qa}.md",
        question=extra.pop("question", f"Question {qa}?"),
        answer=extra.pop("answer", f"Answer {qa}"),
        context=f"Passage of {qa}.",
        status=extra.pop("status", PairStatus.ACTIVE.value),
        model="gen/model",
        question_hash=qa,
        created_at=T0 + timedelta(minutes=minute),
        **extra,
    )


def _stamp(
    stage: str, outcome: str, note: str = "", **more: Any
) -> dict[str, list[dict[str, Any]]]:
    entry = decisions.record(job_id=JOB, stage=stage, outcome=outcome, note=note)
    return {decisions.SCREENING: [{**entry, **more}]}


def _seed(*, answers: bool = True) -> None:
    """A job that wrote, filtered and asked; written out of the question order."""
    writer = {
        "call": "call-1",
        "model": "gen/model",
        "system": "You write questions.",
        "user": f"Passage:\n{PASSAGE_MARK}\nWrite two.",
        "reply": '{"items": []}',
        "cost_usd": 0.002,
        "latency_s": 1.5,
    }
    web = {
        "model": "web/model",
        "judge": J1,
        "system": "Answer with search.",
        "user": "Question w1?",
        "answer": "found it",
        "citations": [{"url": "https://example.org/a", "title": "A"}],
        "verdict": "correct",
        "reasoning": "matches",
        "search": "native",
        "web_search": True,
        "web_search_requests": 2,
        "judge_system": "Grade it.",
        "judge_prompt": "Gold: x",
        "judge_reply": "CORRECT",
        "cost_usd": 0.01,
        "latency_s": 2.0,
        "judge_cost_usd": 0.001,
        "judge_latency_s": 0.5,
    }
    gate = {
        "model": J1,
        "system": "Is it in the fragments?",
        "user": "Question c1? + fragments",
        "reply": "NOT FOUND",
        "cost_usd": 0.003,
        "latency_s": 0.7,
    }
    rows: list[Any] = [
        # Written first, but qa sorts after the masking kinds.
        _pair("qa1", "qa", 0, meta={"writer": writer}),
        _pair("qa2", "qa", 1, meta={"writer": {"call": "call-1"}}),
        _pair(
            "m1",
            "named_entity_masking",
            2,
            meta={"category": "named_entity", "entity_label": "ORG"},
        ),
        _pair(
            "mcq1",
            "mcq",
            3,
            distractors=["B) two", "C) three", "D) four"],
            answer="A) one",
        ),
        _pair(
            "w1",
            "qa",
            4,
            status=PairStatus.REJECTED.value,
            status_reason=StatusReason.WEB_ANSWERABLE.value,
            meta=_stamp(
                decisions.WEB_CHECK,
                decisions.REMOVED,
                "web_answerable",
                reason=StatusReason.WEB_ANSWERABLE.value,
                web=web,
            ),
        ),
        _pair(
            "c1",
            "unanswerable_property",
            5,
            status=PairStatus.REJECTED.value,
            status_reason=StatusReason.RETRIEVAL_GATE.value,
            meta=_stamp(
                decisions.CONTROL,
                decisions.REMOVED,
                "found in the data",
                reason=StatusReason.RETRIEVAL_GATE.value,
                gate=gate,
            ),
        ),
        _pair(
            "cap1",
            "qa",
            6,
            status=PairStatus.RETIRED.value,
            status_reason=StatusReason.ROTATION.value,
            status_note="over the set cap",
        ),
        # An earlier job's question this job re-checked.
        _pair(
            "old1",
            "qa",
            -60,
            job_id=EARLIER,
            meta=_stamp(decisions.WEB_CHECK, decisions.KEPT),
        ),
    ]
    results: list[Any] = []
    if answers:
        n = 0

        def answer(
            qa: str,
            model: str,
            arm: str,
            judge: str,
            verdict: str,
            *,
            block: str = "direct",
            text: str = "an answer",
            extra: dict[str, Any] | None = None,
        ) -> None:
            nonlocal n
            n += 1
            run_id = f"{JOB}-run-{n}"
            rows.append(
                Run(
                    id=run_id,
                    space=KEY,
                    job_id=JOB,
                    context_mode=arm,
                    block=block,
                    model=model,
                    judge_model=judge,
                    params={"context_docs": 3},
                )
            )
            results.append(
                Result(
                    id=f"{JOB}-res-{n}",
                    run_id=run_id,
                    qa_id=qa,
                    space=KEY,
                    answer=text,
                    verdict=verdict,
                    reasoning=f"{judge} says {verdict}",
                    judge_model=judge,
                    model=model,
                    latency_s=1.25,
                    extra=extra or {},
                    audit={
                        "responder_system": "Answer briefly.",
                        "responder_prompt": f"Q: {qa}\n\nExcerpts: ...",
                        "context": "Excerpts: ...",
                        "call": {"finish_reason": "stop", "cost_usd": 0.0004},
                        "judge_system": "Judge it.",
                        "judge_prompt": f"Grade {qa}",
                        "judge_raw": verdict.upper(),
                        "judge_call": {"cost_usd": 0.0001, "latency_s": 0.4},
                    },
                    created_at=T0 + timedelta(hours=1, seconds=n),
                )
            )

        # Inserted in reverse of the expected order.
        answer("qa1", M2, "model_with_context", J1, "correct", text=LONG)
        answer("qa1", M1, "model_with_context", J2, "hallucinate")
        answer("qa1", M1, "model_with_context", J1, "correct")
        answer("qa1", M1, "closed_book", J1, "abstain")
        answer(
            "m1",
            M1,
            "model_with_context",
            J1,
            "correct",
            block="denial_loop",
            extra={
                "denial": {
                    "rounds": 2,
                    "judge_system": "Judge rounds.",
                    "log": [
                        {
                            "round": 1,
                            "objection": "Are you sure?",
                            "answer": "Yes.",
                            "judge_prompt": "Round 1 grade",
                            "judge_raw": "HELD",
                            "cost_usd": 0.0002,
                        },
                        {"round": 2, "objection": "Really?", "answer": "Yes!"},
                    ],
                }
            },
        )
        answer("m1", M1, "model_with_context", J1, "correct")
        answer(
            "mcq1",
            M1,
            "closed_book",
            J1,
            "correct",
            block="monte_carlo",
            extra={
                "monte_carlo": {
                    "log": [
                        {"trial": 1, "temperature": 0.7, "answer": "A", "correct": True}
                    ]
                }
            },
        )

    with session_scope() as session:
        session.add(
            Job(
                id=EARLIER,
                target=KEY,
                state=JobState.SUCCEEDED.value,
                created_at=T0 - timedelta(days=1),
            )
        )
        session.add(
            Job(
                id=JOB,
                target=KEY,
                state=JobState.SUCCEEDED.value,
                created_at=T0,
                started_at=T0 + timedelta(minutes=1),
                finished_at=T0 + timedelta(minutes=55),
                params={
                    "judge_panel": [J1, J2],
                    "web_check_model": "web/model",
                    "web_check_judge": J1,
                    "cost": {
                        "usd": 1.5,
                        "usd_calls": 1.4,
                        "spend_before": 10.0,
                        "spend_after": 11.5,
                        "total_usd": 1.4,
                        "by_role": {
                            "writer": 0.1,
                            "web_check": 0.2,
                            "subjects": 0.9,
                            "judges": 0.2,
                        },
                    },
                    "timing": {
                        "total_s": 3240.0,
                        "phases": [
                            {"phase": "generate", "s": 600.0},
                            {"phase": "evaluate", "s": 2400.0},
                        ],
                        "passes": [
                            {
                                "arm": "closed_book",
                                "block": "direct",
                                "model": M1,
                                "s": 300.0,
                                "questions": 2,
                                "stopped": "",
                            }
                        ],
                        "calls": [
                            {
                                "model": M1,
                                "role": "answer",
                                "count": 5,
                                "failed": 0,
                                "mean_s": 1.0,
                                "p90_s": 2.0,
                                "max_s": 3.0,
                                "total_s": 5.0,
                            }
                        ],
                        "concurrency": {"model": 8, "endpoint": 4},
                    },
                },
            )
        )
        session.add_all([r for r in rows if isinstance(r, QaPair)])
        session.flush()
        session.add_all([r for r in rows if isinstance(r, Run)])
        session.flush()
        session.add_all(results)


def _sheet(wb: Any, name: str) -> list[dict[str, Any]]:
    rows = list(wb[name].iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, row, strict=True)) for row in rows[1:]]


def _download(client: TestClient, auth: dict[str, str], job: str = JOB) -> Any:
    got = client.get(f"/console/report/runs/{job}/export.xlsx", headers=auth)
    assert got.status_code == 200, got.text
    assert got.headers["content-type"] == export_xlsx.MEDIA_TYPE
    assert got.headers["content-disposition"] == (
        'attachment; filename="support-kb-run-2026-09-01-0600.xlsx"'
    )
    return load_workbook(io.BytesIO(got.content))


def test_the_workbook_holds_every_part_in_order(client: TestClient, clean: Any) -> None:
    auth = _console(client)
    _seed()
    wb = _download(client, auth)
    assert wb.sheetnames == list(export_xlsx.SHEETS)

    for name in export_xlsx.SHEETS:
        ws = wb[name]
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref.startswith("A1:")

    run = {row["Field"]: row["Value"] for row in _sheet(wb, "Run")}
    assert run["Endpoint"] == "Support KB"
    assert run["Duration (s)"] == 3240
    assert run["Cost total (USD)"] == 1.4
    assert run["Cost: subjects (USD)"] == 0.9
    assert run["Key spend delta (USD)"] == 1.5
    assert run["Phase: evaluate (s)"] == 2400
    assert run["Concurrency: model"] == 8
    assert run["Web check model"] == "web/model"
    assert run["Written"] == 7
    assert run["Asked"] == 3
    assert run["Settings: context_docs"] == 3

    generated = _sheet(wb, "Generated")
    assert [r["Question id"] for r in generated] == [
        "m1",
        "mcq1",
        "qa1",
        "qa2",
        "w1",
        "cap1",
        "c1",
    ]
    qa1, qa2 = generated[2], generated[3]
    assert qa1["Writer prompt (system)"] == "You write questions."
    assert qa1["Writer prompt (user)"] == "Passage:\nPassage of qa1.\nWrite two."
    # The call's texts sit on qa1 only; qa2 shows them too.
    assert qa2["Writer reply"] == '{"items": []}'
    assert qa2["Writer call cost (USD)"] == 0.002
    assert generated[0]["Writer prompt (system)"] is None
    assert "Blank: named_entity ORG" in generated[0]["Options / key facts / blanks"]
    assert "B) two" in generated[1]["Options / key facts / blanks"]

    rejected = _sheet(wb, "Rejected")
    assert [(r["Question id"], r["Stage"]) for r in rejected] == [
        ("w1", "web_check"),
        ("cap1", "cap"),
        ("c1", "control"),
    ]

    assert [r["Question id"] for r in _sheet(wb, "Kept")] == ["m1", "mcq1", "qa1"]

    flt = {r["Question id"]: r for r in _sheet(wb, "Filter")}
    # Unstamped questions of the job get a decision read off their status.
    assert list(flt) == ["m1", "mcq1", "old1", "qa1", "qa2", "w1", "cap1", "c1"]
    assert flt["qa1"]["Recorded"] == "no"
    old, w1, c1 = flt["old1"], flt["w1"], flt["c1"]
    assert old["Written earlier"] == "yes" and old["Outcome"] == "kept"
    assert w1["Web check prompt (user)"] == "Question w1?"
    assert w1["Judge reply"] == "CORRECT"
    assert w1["Citations"] == "https://example.org/a A"
    assert w1["Searches"] == 2
    assert w1["Cost (USD)"] == pytest.approx(0.011)
    assert w1["Time (s)"] == pytest.approx(2.5)
    assert c1["Control gate reply"] == "NOT FOUND"
    assert c1["Cost (USD)"] == pytest.approx(0.003)

    answers = _sheet(wb, "Answers")
    assert [
        (r["Question id"], r["Model"], r["Condition"], r["Check"]) for r in answers
    ] == [
        ("m1", M1, "model_with_context", "direct"),
        ("m1", M1, "model_with_context", "denial_loop: first answer"),
        ("m1", M1, "model_with_context", "challenge round 1"),
        ("m1", M1, "model_with_context", "challenge round 2"),
        ("mcq1", M1, "closed_book", "monte_carlo: first answer"),
        ("mcq1", M1, "closed_book", "repeat at temperature 0.7, trial 1"),
        ("qa1", M1, "closed_book", "direct"),
        ("qa1", M2, "model_with_context", "direct"),
        ("qa1", M1, "model_with_context", "direct"),
    ]
    assert answers[2]["User prompt"] == "Are you sure?"
    assert answers[2]["For judge"] == J1
    assert answers[5]["User prompt"].startswith("Q: mcq1")
    assert answers[0]["Cost (USD)"] == 0.0004

    # A text past the cell limit continues in part columns of the same row.
    long_row = answers[7]
    parts = [
        long_row["Answer"],
        long_row["Answer (part 2)"],
        long_row["Answer (part 3)"],
    ]
    assert all(len(p) <= export_xlsx.CELL_LIMIT for p in parts)
    assert "".join(parts) == LONG
    assert answers[0]["Answer (part 2)"] is None

    judging = _sheet(wb, "Judging")
    assert [(r["Question id"], r["Judge"], r["Check"]) for r in judging] == [
        ("m1", J1, "direct"),
        ("m1", J1, "denial_loop"),
        ("m1", J1, "challenge round 1"),
        ("mcq1", J1, "monte_carlo"),
        ("qa1", J1, "direct"),
        ("qa1", J1, "direct"),
        ("qa1", J1, "direct"),
        ("qa1", J2, "direct"),
    ]
    counted = [(r["Model"], r["Judge"], r["Counted"]) for r in judging[6:8]]
    assert counted == [(M1, J1, "yes"), (M1, J2, "no")]
    assert judging[2]["Judge prompt (system)"] == "Judge rounds."
    assert judging[0]["Cost (USD)"] == 0.0001

    timing = _sheet(wb, "Timing")
    assert [r["Part"] for r in timing] == ["total", "phase", "phase", "pass", "calls"]


def test_a_build_only_job_and_older_records(client: TestClient, clean: Any) -> None:
    auth = _console(client)
    _seed(answers=False)
    wb = _download(client, auth)
    assert _sheet(wb, "Answers") == []
    assert _sheet(wb, "Judging") == []
    # Nothing asked: kept = what it let into the set.
    assert [r["Question id"] for r in _sheet(wb, "Kept")] == [
        "m1",
        "mcq1",
        "old1",
        "qa1",
        "qa2",
    ]


def test_a_foreign_job_is_a_404(client: TestClient, clean: Any) -> None:
    _console(client)
    other = _console(client, KEY2)
    _seed()
    got = client.get(f"/console/report/runs/{JOB}/export.xlsx", headers=other)
    assert got.status_code == 404


def test_three_thousand_answers_export_quickly(client: TestClient, clean: Any) -> None:
    auth = _console(client)
    _seed(answers=False)
    runs = [
        Run(
            id=f"{JOB}-big-{arm}-{judge}",
            space=KEY,
            job_id=JOB,
            context_mode=arm,
            block="direct",
            model=M1,
            judge_model=judge,
        )
        for arm in ("closed_book", "model_with_context")
        for judge in (J1, J2)
    ]
    pairs = [_pair(f"big{n:04d}", "qa", 100 + n) for n in range(750)]
    prompt = "Excerpt. " * 400
    results = [
        Result(
            id=f"{run.id}-{pair.id}",
            run_id=run.id,
            qa_id=pair.id,
            space=KEY,
            answer="An answer. " * 50,
            verdict="correct",
            judge_model=run.judge_model,
            model=M1,
            audit={
                "responder_system": "Answer briefly.",
                "responder_prompt": prompt,
                "judge_prompt": prompt,
                "judge_raw": "CORRECT",
            },
        )
        for run in runs
        for pair in pairs
    ]
    with session_scope() as session:
        session.add_all(pairs)
        session.add_all(runs)
        session.flush()
        session.add_all(results)

    started = time.monotonic()
    got = client.get(f"/console/report/runs/{JOB}/export.xlsx", headers=auth)
    took = time.monotonic() - started
    assert got.status_code == 200, got.text
    assert took < 30, took
    wb = load_workbook(io.BytesIO(got.content), read_only=True)
    rows = {
        name: sum(1 for _ in wb[name].iter_rows()) for name in ("Judging", "Answers")
    }
    assert rows == {"Judging": 3001, "Answers": 1501}
